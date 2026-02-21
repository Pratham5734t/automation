import os
import logging
from datetime import datetime, timezone
import openpyxl

logger = logging.getLogger(__name__)

HEADERS = ["shortcode", "source_account", "posted_at"]


class PostStorage:
    """Manages post history using an Excel file to avoid duplicate uploads."""

    def __init__(self, filepath: str = "history.xlsx"):
        self.filepath = filepath
        self._posted: set = set()
        self._ensure_file()
        self._load_history()

    def _ensure_file(self) -> None:
        if not os.path.exists(self.filepath):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "History"
            ws.append(HEADERS)
            wb.save(self.filepath)
            logger.info("Created new history file: %s", self.filepath)

    def _load_history(self) -> None:
        wb = openpyxl.load_workbook(self.filepath)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                self._posted.add(row[0])

    def already_posted(self, shortcode: str) -> bool:
        """Return True if *shortcode* has already been uploaded."""
        return shortcode in self._posted

    def mark_posted(self, shortcode: str, source_account: str) -> None:
        """Record *shortcode* as successfully uploaded."""
        wb = openpyxl.load_workbook(self.filepath)
        ws = wb.active
        ws.append([shortcode, source_account, datetime.now(timezone.utc).isoformat()])
        wb.save(self.filepath)
        self._posted.add(shortcode)
        logger.info("Marked %s as posted", shortcode)
